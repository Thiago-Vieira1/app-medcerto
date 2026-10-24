import unicodedata
from collections import OrderedDict
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apimedcerto.models import Medication


class MedicationImportError(Exception):
    pass


def _normalized_header(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char))
    return " ".join(value.strip().upper().split())


def _normalized_name(value):
    return " ".join(str(value or "").strip().upper().split())


def _cell_text(value):
    return "" if value is None else str(value).strip()


def import_medications(uploaded_file):
    if not uploaded_file.name.lower().endswith(".xlsx"):
        raise MedicationImportError("Envie um arquivo Excel .xlsx.")

    workbook = None
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        if "Lista PMC" not in workbook.sheetnames:
            raise MedicationImportError("A planilha precisa conter a aba 'Lista PMC'.")

        worksheet = workbook["Lista PMC"]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise MedicationImportError("A aba 'Lista PMC' está vazia.")

        header_indexes = {
            _normalized_header(header): index
            for index, header in enumerate(headers)
            if header is not None
        }
        required_headers = {
            "PRODUTO": "PRODUTO",
            "TARJA": "TARJA",
            "CLASSE TERAPEUTICA": "CLASSE TERAPÊUTICA",
            "REGISTRO": "REGISTRO",
        }
        missing_headers = [
            display_name
            for normalized, display_name in required_headers.items()
            if normalized not in header_indexes
        ]
        if missing_headers:
            raise MedicationImportError(
                "Colunas obrigatórias ausentes: " + ", ".join(missing_headers) + "."
            )

        products = OrderedDict()
        rows_read = 0
        rows_without_name = 0

        for excel_row, row in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            rows_read += 1
            name = _normalized_name(row[header_indexes["PRODUTO"]])
            if not name:
                rows_without_name += 1
                continue

            stripe = _cell_text(row[header_indexes["TARJA"]])
            registration = _cell_text(row[header_indexes["REGISTRO"]])
            if len(name) > 255 or len(stripe) > 50 or len(registration) > 20:
                raise MedicationImportError(
                    f"Há um valor maior que o limite permitido na linha {excel_row}."
                )

            products.setdefault(
                name,
                Medication(
                    name=name,
                    stripe=stripe,
                    therapeutic_class=_cell_text(
                        row[header_indexes["CLASSE TERAPEUTICA"]]
                    ),
                    anvisa_registration=registration,
                ),
            )
    except MedicationImportError:
        raise
    except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError, KeyError) as exc:
        raise MedicationImportError("Não foi possível ler o arquivo Excel enviado.") from exc
    finally:
        if workbook is not None:
            workbook.close()

    existing_names = {
        _normalized_name(name)
        for name in Medication.objects.values_list("name", flat=True)
    }
    new_products = [
        medication
        for name, medication in products.items()
        if name not in existing_names
    ]

    with transaction.atomic():
        previous_count = Medication.objects.count()
        Medication.objects.bulk_create(new_products, ignore_conflicts=True)
        created = Medication.objects.count() - previous_count

    return {
        "rows_read": rows_read,
        "unique_products": len(products),
        "created": created,
        "already_existing": len(products) - created,
        "duplicate_rows": rows_read - rows_without_name - len(products),
        "rows_without_name": rows_without_name,
    }
